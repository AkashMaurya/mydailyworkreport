"""Change 2 — combined team report (merge by member + export in all formats)."""
from io import BytesIO

import pytest

from tests.conftest import create_entry, create_user
from app import build_team_report, excel_sheet_title


def _seed_team():
    """Two members with entries across different dates."""
    alex = create_user("Alex Member", "alex@x.com")
    blake = create_user("Blake Member", "blake@x.com")
    create_entry(alex["id"], title="Alex task one", date="2026-10-01", duration_minutes=30)
    create_entry(alex["id"], title="Alex task two", date="2026-10-02", duration_minutes=90)
    create_entry(blake["id"], title="Blake task one", date="2026-10-01", duration_minutes=60)
    return alex, blake


def _entries_for(user_id=None):
    from app import fetch_entries_for_export
    conn_filter = str(user_id) if user_id else None
    return fetch_entries_for_export(conn_filter, None, None, None, {"role": "manager", "id": 1})


# ---------- Merge logic ----------

def test_build_team_report_groups_person_first():
    _seed_team()
    entries = _entries_for()
    grouped, totals, counts = build_team_report(entries)

    # person-first structure, alphabetical persons, chronological dates
    assert list(grouped.keys()) == ["Alex Member", "Blake Member"]
    assert list(grouped["Alex Member"].keys()) == ["2026-10-01", "2026-10-02"]
    assert list(grouped["Blake Member"].keys()) == ["2026-10-01"]

    # entries labeled with their member's name
    for pn, dates in grouped.items():
        for d, ents in dates.items():
            for e in ents:
                assert e["person_name"] == pn

    assert counts == {"Alex Member": 2, "Blake Member": 1}
    assert totals == {"Alex Member": 120, "Blake Member": 60}


def test_build_team_report_empty():
    grouped, totals, counts = build_team_report([])
    assert grouped == {}
    assert dict(totals) == {}
    assert dict(counts) == {}


def test_build_team_report_single_person():
    alex = create_user("Alex Member", "alex@x.com")
    create_entry(alex["id"], title="Solo")
    grouped, totals, counts = build_team_report(_entries_for())
    assert list(grouped.keys()) == ["Alex Member"]
    assert counts["Alex Member"] == 1


def test_excel_sheet_title_sanitization():
    used = set()
    assert excel_sheet_title("Bob / Alice", used) == "Bob - Alice"
    long = "A" * 50
    assert len(excel_sheet_title(long, used)) <= 31
    # duplicates get numeric suffixes
    assert excel_sheet_title("Same Name", used) == "Same Name"
    assert excel_sheet_title("Same Name", used) == "Same Name (2)"
    assert excel_sheet_title("Same Name", used) == "Same Name (3)"
    # empty / illegal-only names fall back
    assert excel_sheet_title("[]", used) == "Sheet" or "-" in excel_sheet_title("[]", used)


# ---------- On-page combined report ----------

def test_team_report_page_shows_person_labeled_sections(manager_client):
    _seed_team()
    resp = manager_client.get("/manager/team-report")
    assert resp.status_code == 200
    html = resp.get_data(as_text=True)
    assert "Combined team report" in html
    assert "Alex Member" in html
    assert "Blake Member" in html
    assert "Alex task one" in html
    assert "Blake task one" in html
    # export buttons for all three formats
    assert 'href="/export/team/excel' in html
    assert 'href="/export/team/word' in html
    assert 'href="/export/team/pdf' in html
    # entry cards label authorship
    assert "By Alex Member" in html


def test_team_report_summary_totals(manager_client):
    _seed_team()
    html = manager_client.get("/manager/team-report").get_data(as_text=True)
    assert "Summary" in html
    assert "Grand total" in html


def test_team_report_respects_filters(manager_client):
    alex, blake = _seed_team()
    resp = manager_client.get("/manager/team-report?person=%d" % blake["id"])
    html = resp.get_data(as_text=True)
    assert "Blake task one" in html
    assert "Alex task one" not in html


def test_team_report_empty_state(manager_client):
    resp = manager_client.get("/manager/team-report")
    assert resp.status_code == 200
    assert "No entries found for the selected filters." in resp.get_data(as_text=True)


def test_team_report_member_forbidden(member_client):
    _seed_team()
    assert member_client.get("/manager/team-report").status_code == 403
    assert member_client.get("/export/team/excel").status_code == 403
    assert member_client.get("/export/team/word").status_code == 403
    assert member_client.get("/export/team/pdf").status_code == 403


def test_team_report_requires_login(client):
    assert client.get("/manager/team-report").status_code == 302


def test_manager_dashboard_links_to_combined_report(manager_client):
    html = manager_client.get("/manager").get_data(as_text=True)
    assert 'href="/manager/team-report' in html


# ---------- Excel export ----------

def test_excel_export_structure(manager_client):
    from openpyxl import load_workbook
    _seed_team()
    resp = manager_client.get("/export/team/excel")
    assert resp.status_code == 200
    assert resp.headers["Content-Disposition"].endswith('.xlsx"') or ".xlsx" in resp.headers["Content-Disposition"]

    wb = load_workbook(BytesIO(resp.data))
    assert "Summary" in wb.sheetnames
    assert "Alex Member" in wb.sheetnames
    assert "Blake Member" in wb.sheetnames

    # Summary sheet: names, counts, totals
    ws = wb["Summary"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    assert rows[0] == ["Team Member", "Entries", "Total Time (min)", "Total Time"]
    assert ["Alex Member", 2, 120, "2h"] in rows
    assert ["Blake Member", 1, 60, "1h"] in rows
    assert ["Grand total", 3, 180, "3h"] in rows

    # Per-person sheet is clearly labeled and holds only that person's entries
    ws = wb["Alex Member"]
    assert ws["A1"].value == "Alex Member"
    body = [str(c) for row in ws.iter_rows(values_only=True) for c in row if c]
    assert "Alex task one" in body
    assert "Alex task two" in body
    assert "Blake task one" not in body


def test_excel_export_empty(manager_client):
    from openpyxl import load_workbook
    resp = manager_client.get("/export/team/excel")
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.data))
    assert "Summary" in wb.sheetnames
    rows = [list(r) for r in wb["Summary"].iter_rows(values_only=True)]
    flat = [str(c) for row in rows for c in row]
    assert "No entries found for the selected filters." in flat


# ---------- Word export ----------

def _docx_paragraphs(data):
    from docx import Document
    doc = Document(BytesIO(data))
    return [p for p in doc.paragraphs]


def test_word_export_person_first_sections(manager_client):
    _seed_team()
    resp = manager_client.get("/export/team/word")
    assert resp.status_code == 200
    assert ".docx" in resp.headers["Content-Disposition"]

    paras = _docx_paragraphs(resp.data)
    texts = [(p.text, p.style.name) for p in paras]

    # Title
    assert texts[0][0] == "Combined Team Work Report"

    # Members appear as top-level (Heading 1) sections
    h1s = [t for t, s in texts if s == "Heading 1"]
    assert "Alex Member" in h1s
    assert "Blake Member" in h1s
    # alphabetical person order
    assert h1s.index("Alex Member") < h1s.index("Blake Member")

    # Entries under the correct person: Alex's heading precedes Alex's tasks
    all_texts = [t for t, _ in texts]
    assert all_texts.index("Alex Member") < all_texts.index("Alex task one")
    assert all_texts.index("Blake Member") < all_texts.index("Blake task one")
    # and Alex's entries do not appear under Blake's section
    assert all_texts.index("Alex task two") < all_texts.index("Blake Member")


def test_word_export_empty(manager_client):
    resp = manager_client.get("/export/team/word")
    assert resp.status_code == 200
    all_texts = [p.text for p in _docx_paragraphs(resp.data)]
    assert "No entries found for the selected filters." in all_texts


# ---------- PDF export ----------

def test_pdf_export(manager_client):
    _seed_team()
    resp = manager_client.get("/export/team/pdf")
    assert resp.status_code == 200
    assert resp.headers["Content-Type"] == "application/pdf"
    assert resp.data[:4] == b"%PDF"
    assert ".pdf" in resp.headers["Content-Disposition"]


def test_pdf_export_empty_still_valid_pdf(manager_client):
    resp = manager_client.get("/export/team/pdf")
    assert resp.status_code == 200
    assert resp.data[:4] == b"%PDF"


def test_pdf_html_template_groups_by_person():
    """The pdfkit source template renders person-first sections."""
    from app import app as flask_app
    _seed_team()
    entries = _entries_for()
    grouped, totals, counts = build_team_report(entries)
    from datetime import datetime
    with flask_app.test_request_context("/manager/team-report"):
        html = flask_app.jinja_env.get_template("report_team_pdf.html").render(
            entries=entries, grouped=grouped, totals=totals, counts=counts,
            date_from=None, date_to=None,
            generated=datetime.now().strftime("%B %d, %Y"),
            requester={"name": "Mara", "role": "manager"},
        )
    # person sections precede their dates/entries
    assert html.index("Alex Member") < html.index("Alex task one")
    assert html.index("Blake Member") < html.index("Blake task one")
    assert "person-section" in html
